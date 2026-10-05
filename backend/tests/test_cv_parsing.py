from __future__ import annotations

import io
import zipfile
from datetime import date
from decimal import Decimal

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas

from app.ai import pdf_docx_extractor
from app.ai.education_extractor import extract_educations
from app.ai.entity_recognizer import extract_profile
from app.ai.errors import ResumeParseError, ResumeParseErrorKind
from app.ai.experience_extractor import extract_experiences
from app.ai.pdf_docx_extractor import (
    DOCX_MIME_TYPE,
    PDF_MIME_TYPE,
    extract_docx_text,
    extract_pdf_text,
    extract_resume_text,
)
from app.ai.resume_parser import RESUME_TEXT_PREPROCESSING_VERSION, parse_resume
from app.ai.schemas import ParsedResumeResult, TaxonomySkill
from app.ai.sections import ResumeSection, identify_section_heading, split_resume_sections
from app.ai.skill_normalizer import normalize_skills
from app.ai.text_processing import MAX_EXTRACTED_TEXT_CHARS, normalize_extracted_text
from app.core.exceptions import APIError
from app.services.resume_service import validate_resume_upload

STRICT_WORDPROCESSINGML_NAMESPACE = "http://purl.oclc.org/ooxml/wordprocessingml/main"
TRANSITIONAL_WORDPROCESSINGML_NAMESPACE = (
    "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
)
STRICT_OFFICE_DOCUMENT_RELATIONSHIP = (
    "http://purl.oclc.org/ooxml/officeDocument/relationships/officeDocument"
)
TRANSITIONAL_OFFICE_DOCUMENT_RELATIONSHIP = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument"
)


def make_pdf(pages: list[list[str]]) -> bytes:
    output = io.BytesIO()
    pdf = canvas.Canvas(output)
    for lines in pages:
        y = 800
        for line in lines:
            pdf.drawString(72, y, line)
            y -= 18
        pdf.showPage()
    pdf.save()
    return output.getvalue()


def make_docx(paragraphs: list[str], tables: list[list[list[str]]] | None = None) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    for table_data in tables or []:
        table = document.add_table(rows=len(table_data), cols=len(table_data[0]))
        for row_index, row in enumerate(table_data):
            for column_index, value in enumerate(row):
                table.cell(row_index, column_index).text = value
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def make_profile_docx(
    document_xml: str | bytes,
    relationship_type: str = STRICT_OFFICE_DOCUMENT_RELATIONSHIP,
    *,
    target_mode: str | None = None,
    additional_relationships: tuple[str, ...] = (),
) -> bytes:
    target_mode_attribute = f" TargetMode='{target_mode}'" if target_mode is not None else ""
    primary_relationship = (
        f"<Relationship Id='rId1' Type='{relationship_type}' "
        f"Target='word/document.xml'{target_mode_attribute}/>"
    )
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            "<Types xmlns='http://schemas.openxmlformats.org/package/2006/content-types'>"
            "<Override PartName='/word/document.xml' ContentType='application/vnd."
            "openxmlformats-officedocument.wordprocessingml.document.main+xml'/>"
            "</Types>",
        )
        archive.writestr(
            "_rels/.rels",
            "<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/"
            "relationships'>"
            + primary_relationship
            + "".join(additional_relationships)
            + "</Relationships>",
        )
        archive.writestr("word/document.xml", document_xml)
    return output.getvalue()


def strict_document(body: str) -> str:
    return (
        f"<w:document xmlns:w='{STRICT_WORDPROCESSINGML_NAMESPACE}'>"
        f"<w:body>{body}</w:body></w:document>"
    )


@pytest.fixture
def taxonomy() -> tuple[TaxonomySkill, ...]:
    names = [
        (1, "Java", "java", "HARD", "Language"),
        (2, "JavaScript", "javascript", "HARD", "Language"),
        (3, "C#", "csharp", "HARD", "Language"),
        (4, "C++", "cpp", "HARD", "Language"),
        (5, "React", "react", "HARD", "Frontend"),
        (6, "Next.js", "nextjs", "HARD", "Frontend"),
        (7, "Node.js", "nodejs", "HARD", "Backend"),
        (8, "ASP.NET Core", "aspnet-core", "HARD", "Backend"),
        (9, "CI/CD", "cicd", "HARD", "DevOps"),
        (10, "Tailwind CSS", "tailwind-css", "HARD", "Frontend"),
        (11, "Spring Boot", "spring-boot", "HARD", "Backend"),
        (12, "Communication", "communication", "SOFT", "Communication"),
        (13, "Python", "python", "HARD", "Language"),
    ]
    return tuple(
        TaxonomySkill(
            id=skill_id,
            name=name,
            normalized_name=normalized,
            skill_kind=kind,  # type: ignore[arg-type]
            category=category,
        )
        for skill_id, name, normalized, kind, category in names
    )


def test_real_pdf_extracts_text_and_preserves_page_order() -> None:
    source = make_pdf([["First page", "Python"], ["Second page", "JavaScript"]])
    text = extract_pdf_text(source)
    assert text.index("First page") < text.index("Second page")
    assert "Python" in text and "JavaScript" in text


def test_real_docx_extracts_paragraphs_and_table_cells_in_stable_order() -> None:
    source = make_docx(["Header paragraph"], [[["Left cell", "Right cell"]]])
    text = extract_docx_text(source)
    assert text.splitlines() == ["Header paragraph", "Left cell", "Right cell"]


def test_strict_docx_paragraph_is_accepted_by_upload_and_parser() -> None:
    source = make_profile_docx(strict_document("<w:p><w:r><w:t>Strict resume</w:t></w:r></w:p>"))

    _, mime_type = validate_resume_upload("resume.docx", source)

    assert mime_type == DOCX_MIME_TYPE
    assert extract_docx_text(source) == "Strict resume"


def test_strict_docx_table_text_preserves_document_order() -> None:
    source = make_profile_docx(
        strict_document(
            "<w:p><w:r><w:t>Before table</w:t></w:r></w:p>"
            "<w:tbl><w:tr>"
            "<w:tc><w:p><w:r><w:t>Left cell</w:t></w:r></w:p></w:tc>"
            "<w:tc><w:p><w:r><w:t>Right cell</w:t></w:r></w:p></w:tc>"
            "</w:tr></w:tbl>"
            "<w:p><w:r><w:t>After table</w:t></w:r></w:p>"
        )
    )

    assert extract_docx_text(source).splitlines() == [
        "Before table",
        "Left cell",
        "Right cell",
        "After table",
    ]


def test_strict_docx_preserves_vietnamese_and_is_deterministic(
    taxonomy: tuple[TaxonomySkill, ...],
) -> None:
    source = make_profile_docx(
        strict_document(
            "<w:p><w:r><w:t>K\u1ef9 n\u0103ng</w:t><w:tab/><w:t>Python</w:t></w:r></w:p>"
            "<w:p><w:r><w:t>Kinh nghi\u1ec7m l\u00e0m vi\u1ec7c</w:t></w:r></w:p>"
        )
    )

    first_text = extract_docx_text(source)
    second_text = extract_docx_text(source)

    assert first_text == second_text
    assert "K\u1ef9 n\u0103ng Python" in first_text
    assert "Kinh nghi\u1ec7m l\u00e0m vi\u1ec7c" in first_text
    assert parse_resume(source, DOCX_MIME_TYPE, taxonomy) == parse_resume(
        source,
        DOCX_MIME_TYPE,
        taxonomy,
    )


@pytest.mark.parametrize(
    ("relationship_type", "document_namespace"),
    [
        (STRICT_OFFICE_DOCUMENT_RELATIONSHIP, TRANSITIONAL_WORDPROCESSINGML_NAMESPACE),
        (TRANSITIONAL_OFFICE_DOCUMENT_RELATIONSHIP, STRICT_WORDPROCESSINGML_NAMESPACE),
    ],
)
def test_docx_mixed_relationship_and_document_profiles_are_rejected(
    relationship_type: str,
    document_namespace: str,
) -> None:
    source = make_profile_docx(
        f"<w:document xmlns:w='{document_namespace}'><w:body/></w:document>",
        relationship_type,
    )

    with pytest.raises(ResumeParseError) as raised:
        extract_docx_text(source)

    assert raised.value.kind is ResumeParseErrorKind.UNREADABLE_SOURCE
    with pytest.raises(APIError) as upload_error:
        validate_resume_upload("resume.docx", source)
    assert upload_error.value.status_code == 415


@pytest.mark.parametrize(
    "source",
    [
        make_profile_docx(
            strict_document("<w:p><w:r><w:t>Strict</w:t></w:r></w:p>"),
            target_mode="External",
        ),
        make_profile_docx(
            f"<w:document xmlns:w='{TRANSITIONAL_WORDPROCESSINGML_NAMESPACE}'>"
            "<w:body><w:p><w:r><w:t>Transitional</w:t></w:r></w:p></w:body>"
            "</w:document>",
            TRANSITIONAL_OFFICE_DOCUMENT_RELATIONSHIP,
            target_mode="External",
        ),
        make_profile_docx(
            strict_document("<w:p><w:r><w:t>Strict</w:t></w:r></w:p>"),
            additional_relationships=(
                "<Relationship Id='rId2' "
                f"Type='{TRANSITIONAL_OFFICE_DOCUMENT_RELATIONSHIP}' "
                "Target='word/document.xml'/>",
            ),
        ),
        make_profile_docx(
            strict_document("<w:p><w:r><w:t>Strict</w:t></w:r></w:p>"),
            "urn:unsupported-office-document-relationship",
        ),
        make_profile_docx(
            strict_document("<w:p><w:r><w:t>Strict</w:t></w:r></w:p>"),
            additional_relationships=("<Relationship",),
        ),
    ],
    ids=["strict-external", "transitional-external", "mixed", "unsupported", "malformed"],
)
def test_upload_and_parser_reject_invalid_main_document_relationships(source: bytes) -> None:
    with pytest.raises(APIError) as upload_error:
        validate_resume_upload("resume.docx", source)
    with pytest.raises(ResumeParseError) as parser_error:
        extract_docx_text(source)

    assert upload_error.value.status_code == 415
    assert parser_error.value.kind is ResumeParseErrorKind.UNREADABLE_SOURCE


def test_malformed_strict_document_is_a_controlled_failure() -> None:
    source = make_profile_docx(
        f"<w:document xmlns:w='{STRICT_WORDPROCESSINGML_NAMESPACE}'>"
        "<w:body><w:p><w:r><w:t>Truncated"
    )

    with pytest.raises(ResumeParseError) as raised:
        extract_docx_text(source)

    assert raised.value.kind is ResumeParseErrorKind.UNREADABLE_SOURCE


def test_strict_document_xml_size_limit_is_controlled(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    document_xml = strict_document("<w:p><w:r><w:t>Resume</w:t></w:r></w:p>")
    source = make_profile_docx(document_xml)
    monkeypatch.setattr(
        pdf_docx_extractor,
        "MAX_DOCX_DOCUMENT_XML_BYTES",
        len(document_xml.encode()) - 1,
    )

    with pytest.raises(ResumeParseError) as raised:
        extract_docx_text(source)

    assert raised.value.kind is ResumeParseErrorKind.TEXT_LIMIT_EXCEEDED


def test_docx_package_entry_limit_rejects_before_python_docx(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = make_docx(["Resume text"])
    document_called = False

    def fail_if_called(_source: io.BytesIO) -> None:
        nonlocal document_called
        document_called = True
        raise AssertionError("python-docx must not run after preflight rejection")

    monkeypatch.setattr(pdf_docx_extractor, "MAX_DOCX_ZIP_ENTRIES", 1)
    monkeypatch.setattr(pdf_docx_extractor, "Document", fail_if_called)

    with pytest.raises(ResumeParseError) as raised:
        extract_docx_text(source)

    assert raised.value.kind is ResumeParseErrorKind.TEXT_LIMIT_EXCEEDED
    assert document_called is False


def test_docx_package_uncompressed_size_limit_is_controlled(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = make_docx(["Resume text"])
    with zipfile.ZipFile(io.BytesIO(source)) as archive:
        declared_size = sum(entry.file_size for entry in archive.infolist())
    monkeypatch.setattr(
        pdf_docx_extractor,
        "MAX_DOCX_UNCOMPRESSED_PACKAGE_BYTES",
        declared_size - 1,
    )

    with pytest.raises(ResumeParseError) as raised:
        extract_docx_text(source)

    assert raised.value.kind is ResumeParseErrorKind.TEXT_LIMIT_EXCEEDED


@pytest.mark.parametrize(
    ("source", "mime_type"),
    [(b"not-pdf", PDF_MIME_TYPE), (b"not-docx", DOCX_MIME_TYPE)],
)
def test_corrupted_documents_raise_controlled_failure(source: bytes, mime_type: str) -> None:
    with pytest.raises(ResumeParseError) as raised:
        extract_resume_text(source, mime_type)
    assert raised.value.kind is ResumeParseErrorKind.UNREADABLE_SOURCE


def test_encrypted_pdf_and_empty_pdf_fail_cleanly() -> None:
    reader = PdfReader(io.BytesIO(make_pdf([["Secret"]])))
    writer = PdfWriter()
    writer.append_pages_from_reader(reader)
    writer.encrypt("password")
    encrypted = io.BytesIO()
    writer.write(encrypted)
    with pytest.raises(ResumeParseError) as encrypted_error:
        extract_pdf_text(encrypted.getvalue())
    assert encrypted_error.value.kind is ResumeParseErrorKind.ENCRYPTED_PDF

    with pytest.raises(ResumeParseError) as empty_error:
        extract_pdf_text(make_pdf([[]]))
    assert empty_error.value.kind is ResumeParseErrorKind.NO_MEANINGFUL_TEXT

    with pytest.raises(ResumeParseError) as empty_docx_error:
        extract_docx_text(make_docx([]))
    assert empty_docx_error.value.kind is ResumeParseErrorKind.NO_MEANINGFUL_TEXT


def test_text_normalization_is_deterministic_and_preserves_vietnamese() -> None:
    source = "  Nguye\u0302\u0303n\t Va\u0306n  An\r\n\x00Kỹ   năng\r\n\r\n\r\nPython  "
    expected = "Nguyễn Văn An\nKỹ năng\n\nPython"
    assert normalize_extracted_text(source) == expected
    assert normalize_extracted_text(source) == normalize_extracted_text(source)


def test_extracted_text_size_guard_is_controlled() -> None:
    with pytest.raises(ResumeParseError) as raised:
        normalize_extracted_text("x" * (MAX_EXTRACTED_TEXT_CHARS + 1))
    assert raised.value.kind is ResumeParseErrorKind.TEXT_LIMIT_EXCEEDED


@pytest.mark.parametrize(
    ("heading", "section"),
    [
        ("professional SUMMARY:", ResumeSection.PROFILE),
        ("KỸ NĂNG CHUYÊN MÔN", ResumeSection.SKILLS),
        ("Kinh nghiệm làm việc", ResumeSection.EXPERIENCE),
        ("hỌc VấN", ResumeSection.EDUCATION),
    ],
)
def test_english_vietnamese_and_mixed_case_section_headings(
    heading: str,
    section: ResumeSection,
) -> None:
    assert identify_section_heading(heading) is section


def test_uppercase_nonheading_is_not_misclassified() -> None:
    assert identify_section_heading("SENIOR SOFTWARE ENGINEER") is None
    sections = split_resume_sections("NGUYỄN VĂN AN\nSENIOR SOFTWARE ENGINEER")
    assert sections.sections == {}


def test_english_sections_are_split_without_uppercase_guessing() -> None:
    sections = split_resume_sections(
        "Candidate Name\nSkills\nPython\nWork Experience\nEngineer | Acme\n"
        "Education\nHanoi University"
    )
    assert sections.get(ResumeSection.SKILLS) == "Python"
    assert sections.get(ResumeSection.EXPERIENCE) == "Engineer | Acme"
    assert sections.get(ResumeSection.EDUCATION) == "Hanoi University"


def test_job_title_is_not_fabricated_as_full_name() -> None:
    text = "SENIOR SOFTWARE ENGINEER\nemail@example.com"
    assert extract_profile(text, split_resume_sections(text)).full_name is None


def test_profile_extracts_evidenced_values_and_explicit_summary() -> None:
    text = (
        "Nguyễn Văn An\n"
        "Title: Backend Engineer\n"
        "Location: Hà Nội\n"
        "an.nguyen@example.com | +84 912 345 678\n"
        "linkedin.com/in/nguyenvanan github.com/nguyenvanan\n\n"
        "Tóm tắt\n"
        "Kỹ sư phần mềm tập trung vào hệ thống backend."
    )
    sections = split_resume_sections(text)
    profile = extract_profile(text, sections)
    assert profile.full_name == "Nguyễn Văn An"
    assert profile.email == "an.nguyen@example.com"
    assert profile.phone_number == "+84 912 345 678"
    assert profile.linkedin_url == "linkedin.com/in/nguyenvanan"
    assert profile.github_url == "github.com/nguyenvanan"
    assert profile.current_title == "Backend Engineer"
    assert profile.location == "Hà Nội"
    assert profile.professional_summary == "Kỹ sư phần mềm tập trung vào hệ thống backend."


def test_profile_missing_values_remain_none_and_years_are_not_phones() -> None:
    text = "Skills\nPython\nExperience\n2019 - 2022"
    profile = extract_profile(text, split_resume_sections(text))
    assert profile.full_name is None
    assert profile.email is None
    assert profile.phone_number is None
    assert profile.current_title is None
    assert profile.location is None
    assert profile.professional_summary is None


def test_skill_matching_handles_punctuation_soft_skills_and_deduplication(
    taxonomy: tuple[TaxonomySkill, ...],
) -> None:
    before = tuple(taxonomy)
    text = (
        "C#, C++, Next.js, Node.js, ASP.NET Core, CI/CD, Tailwind CSS, Spring Boot, "
        "Communication, Python and Python. UnknownQuantumTool."
    )
    skills = normalize_skills(text, taxonomy)
    assert [skill.name for skill in skills] == [
        "C#",
        "C++",
        "Next.js",
        "Node.js",
        "ASP.NET Core",
        "CI/CD",
        "Tailwind CSS",
        "Spring Boot",
        "Communication",
        "Python",
    ]
    assert all(skill.years_of_experience is None for skill in skills)
    assert all(skill.proficiency_level is None for skill in skills)
    assert taxonomy == before


def test_skill_boundaries_prevent_java_and_react_false_positives(
    taxonomy: tuple[TaxonomySkill, ...],
) -> None:
    skills = normalize_skills("JavaScript and reactive programming", taxonomy)
    assert [skill.name for skill in skills] == ["JavaScript"]


def test_english_and_vietnamese_experience_blocks_are_conservative() -> None:
    text = (
        "Software Engineer | Acme Corp\n"
        "01/2020 - Present\n"
        "Built APIs\n\n"
        "Kỹ sư phần mềm | Công ty ABC\n"
        "March 2018 - December 2019\n"
        "Phát triển dịch vụ"
    )
    results = extract_experiences(text)
    assert results[0].job_title == "Software Engineer"
    assert results[0].company_name == "Acme Corp"
    assert results[0].start_date == date(2020, 1, 1)
    assert results[0].end_date is None and results[0].is_current is True
    assert results[1].start_date == date(2018, 3, 1)
    assert results[1].end_date == date(2019, 12, 1)
    assert results[1].is_current is False


def test_experience_missing_or_ambiguous_dates_are_not_invented() -> None:
    results = extract_experiences(
        "Engineer | Evidence Corp\n2019 - ?\nWork description\n\nUnstructured employer line"
    )
    assert len(results) == 1
    assert results[0].start_date is None and results[0].end_date is None
    assert results[0].is_current is False


def test_year_only_experience_range_does_not_fabricate_month_or_day() -> None:
    result = extract_experiences("Engineer | Evidence Corp\n2019 - 2022")[0]
    assert result.start_date is None and result.end_date is None


def test_invalid_experience_date_order_never_emits_invalid_payload() -> None:
    result = extract_experiences("Engineer | Acme\n12/2022 - 01/2020")[0]
    assert result.start_date is None and result.end_date is None


def test_education_extracts_only_explicit_values() -> None:
    result = extract_educations(
        "Hanoi University of Science\n"
        "Degree: Bachelor of Science\n"
        "Field: Computer Science\n"
        "2018 - 2022\n"
        "GPA: 3.6/4.0"
    )[0]
    assert result.institution_name == "Hanoi University of Science"
    assert result.degree == "Bachelor of Science"
    assert result.field_of_study == "Computer Science"
    assert result.start_year == 2018 and result.graduation_year == 2022
    assert result.gpa == Decimal("3.6")


def test_education_missing_options_and_invalid_year_order_are_conservative() -> None:
    results = extract_educations(
        "Trường Đại học Bách khoa\n2023 - 2020\n\nUnknown learning provider\n2018 - 2020"
    )
    assert len(results) == 1
    result = results[0]
    assert result.degree is None and result.field_of_study is None and result.gpa is None
    assert result.start_year is None and result.graduation_year is None


def test_unsupported_mime_type_is_a_controlled_failure() -> None:
    with pytest.raises(ResumeParseError) as raised:
        extract_resume_text(b"resume", "text/plain")
    assert raised.value.kind is ResumeParseErrorKind.UNSUPPORTED_MIME_TYPE


def test_full_english_resume_is_stable_and_structured(
    taxonomy: tuple[TaxonomySkill, ...],
) -> None:
    source = make_docx(
        [
            "Alice Nguyen",
            "Title: Backend Engineer",
            "Location: Hanoi",
            "alice@example.com | 0912 345 678",
            "Summary",
            "Backend engineer building reliable APIs.",
            "Skills",
            "Python, C#, Communication",
            "Experience",
            "Backend Engineer | Acme Corp",
            "01/2021 - Present",
            "Built APIs",
            "",
            "Education",
            "Hanoi University",
            "Degree: Bachelor of Science",
            "Field: Computer Science",
            "2016 - 2020",
        ]
    )
    first = parse_resume(source, DOCX_MIME_TYPE, taxonomy)
    second = parse_resume(source, DOCX_MIME_TYPE, taxonomy)
    assert first == second
    assert isinstance(first, ParsedResumeResult)
    assert first.profile.full_name == "Alice Nguyen"
    assert [skill.name for skill in first.skills] == ["C#", "Communication", "Python"]
    assert first.experiences[0].company_name == "Acme Corp"
    assert first.educations[0].institution_name == "Hanoi University"
    assert RESUME_TEXT_PREPROCESSING_VERSION == "resume-text-v1"


def test_full_vietnamese_resume_preserves_text_and_missing_values(
    taxonomy: tuple[TaxonomySkill, ...],
) -> None:
    source = make_docx(
        [
            "Nguyễn Văn An",
            "Địa điểm: Đà Nẵng",
            "Tóm tắt",
            "Kỹ sư phần mềm chú trọng chất lượng.",
            "Kỹ năng",
            "JavaScript, React",
            "Kinh nghiệm làm việc",
            "Kỹ sư phần mềm | Công ty Việt",
            "03/2022 - Hiện tại",
            "",
            "Học vấn",
            "Trường Đại học Bách khoa",
            "Cử nhân ngành Công nghệ thông tin",
            "2017 - 2021",
        ]
    )
    result = parse_resume(source, DOCX_MIME_TYPE, taxonomy)
    assert "Nguyễn Văn An" in result.raw_text
    assert result.profile.location == "Đà Nẵng"
    assert result.profile.email is None and result.profile.phone_number is None
    assert [skill.name for skill in result.skills] == ["JavaScript", "React"]
    assert result.experiences[0].is_current is True
    assert result.educations[0].field_of_study == "Công nghệ thông tin"
