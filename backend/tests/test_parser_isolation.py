"""Real subprocess regression, bounded IPC, timeout, and parser compatibility."""

from __future__ import annotations

import asyncio
import io
import sys
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from reportlab.pdfgen import canvas

from app.ai.errors import ResumeParseError, ResumeParseErrorKind
from app.ai.job_parser import JobParseError, JobParseErrorKind, parse_job_description
from app.ai.parser_isolation import (
    MAX_REQUEST_BYTES,
    ParserIsolationError,
    ParserIsolationTimeout,
    _exchange,
    _parser_environment,
    run_job_parser_isolated,
    run_resume_parser_isolated,
)
from app.ai.pdf_docx_extractor import DOCX_MIME_TYPE, PDF_MIME_TYPE
from app.ai.resume_parser import parse_resume
from app.ai.schemas import TaxonomySkill


@pytest.fixture
def taxonomy() -> tuple[TaxonomySkill, ...]:
    return (TaxonomySkill(1, "Python", "python", "HARD", "Language"),)


def _make_docx() -> bytes:
    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph("Nguyễn Văn An")
    document.add_paragraph("Kỹ năng")
    document.add_paragraph("Python")
    document.add_paragraph("Kinh nghiệm làm việc")
    document.add_paragraph("Công ty ABC 2022 - 2024")
    document.save(buffer)
    return buffer.getvalue()


def _make_pdf() -> bytes:
    buffer = io.BytesIO()
    document = canvas.Canvas(buffer)
    document.drawString(60, 750, "Resume Python developer")
    document.save()
    return buffer.getvalue()


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("source", "mime"),
    [
        (_make_pdf(), PDF_MIME_TYPE),
        (_make_docx(), DOCX_MIME_TYPE),
    ],
)
async def test_real_isolated_resume_output_matches_baseline(
    source: bytes, mime: str, taxonomy: tuple[TaxonomySkill, ...]
) -> None:
    expected = parse_resume(source, mime, taxonomy)
    actual = await run_resume_parser_isolated(source, mime, taxonomy)
    assert actual == expected
    assert await run_resume_parser_isolated(source, mime, taxonomy) == actual
    if mime == DOCX_MIME_TYPE:
        assert "Nguyễn Văn An" in actual.raw_text


@pytest.mark.asyncio
async def test_real_isolated_job_matches_original_parser(
    taxonomy: tuple[TaxonomySkill, ...],
) -> None:
    raw = "Yêu cầu\nPython bắt buộc\nTối thiểu 2 năm kinh nghiệm\nBằng cử nhân"
    expected = parse_job_description(raw, taxonomy)
    actual = await run_job_parser_isolated(raw, taxonomy)
    assert actual == expected
    assert await run_job_parser_isolated(raw, taxonomy) == actual


@pytest.mark.asyncio
async def test_child_propagates_controlled_parser_error(
    taxonomy: tuple[TaxonomySkill, ...],
) -> None:
    with pytest.raises(ResumeParseError) as captured:
        await run_resume_parser_isolated(b"not a pdf", PDF_MIME_TYPE, taxonomy)
    assert captured.value.kind is ResumeParseErrorKind.UNREADABLE_SOURCE


@pytest.mark.asyncio
async def test_real_child_is_reaped_on_controlled_timeout(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    created: list[asyncio.subprocess.Process] = []
    real_spawn = asyncio.create_subprocess_exec

    async def tracked_spawn(*args: Any, **kwargs: Any) -> asyncio.subprocess.Process:
        process = await real_spawn(*args, **kwargs)
        created.append(process)
        return process

    monkeypatch.setattr(asyncio, "create_subprocess_exec", tracked_spawn)
    with pytest.raises(ParserIsolationTimeout):
        await _exchange(
            {"version": 1, "kind": "job", "raw_content": "Python", "taxonomy": []},
            timeout_seconds=0.3,
            command=(sys.executable, "-c", "import time; time.sleep(5)"),
        )
    assert len(created) == 1
    assert created[0].returncode is not None


@pytest.mark.asyncio
async def test_oversized_request_rejected_before_process_spawn(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    async def must_not_spawn(*args: Any, **kwargs: Any) -> None:
        raise AssertionError("Oversized IPC should not launch a process")

    monkeypatch.setattr(asyncio, "create_subprocess_exec", must_not_spawn)
    with pytest.raises(ParserIsolationError, match="IPC size"):
        await _exchange({"oversized": "a" * MAX_REQUEST_BYTES})


def test_parser_child_environment_does_not_inherit_secrets(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("DATABASE_URL", "postgresql+asyncpg://secret")
    monkeypatch.setenv("CELERY_BROKER_URL", "redis://secret")
    monkeypatch.setenv("JWT_SECRET_KEY", "private")
    environment = _parser_environment()
    assert "DATABASE_URL" not in environment
    assert "CELERY_BROKER_URL" not in environment
    assert "JWT_SECRET_KEY" not in environment
    assert Path(environment["PYTHONPATH"]).name == "backend"
    assert environment["PYTHONNOUSERSITE"] == "1"


@pytest.mark.asyncio
async def test_unsupported_mime_preserves_original_error_kind(
    taxonomy: tuple[TaxonomySkill, ...],
) -> None:
    with pytest.raises(ResumeParseError) as caught:
        await run_resume_parser_isolated(b"input", "text/plain", taxonomy)
    assert caught.value.kind is ResumeParseErrorKind.UNSUPPORTED_MIME_TYPE


@pytest.mark.asyncio
async def test_invalid_job_unicode_preserves_controlled_error_kind(
    taxonomy: tuple[TaxonomySkill, ...],
) -> None:
    with pytest.raises(JobParseError) as captured:
        await run_job_parser_isolated("Job\ud800 content", taxonomy)
    assert captured.value.kind is JobParseErrorKind.INVALID_INPUT
